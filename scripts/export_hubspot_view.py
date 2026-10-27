"""Write the HubSpot-shaped secondary view (Companies, Owners, Tasks) to an .xlsx for Google Sheets import."""
import json, sys, os, datetime, openpyxl
from openpyxl.styles import Font, PatternFill
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'router'))
import hubspot as H, core
def main(accounts_path, out_path):
    A = [json.loads(l) for l in open(accounts_path)]; R = json.load(open(os.path.join(core.ROOT, 'data/sdrs.SYNTHETIC.json')))['sdrs']
    A.sort(key=lambda a: (a['tier'], a['domain'])); now = datetime.datetime.now(datetime.timezone.utc)
    co = [['segr_domain_key', 'name', 'segr_segment', 'segr_tier', 'segr_fit_score', 'segr_intent_score', 'segr_needs_human', 'segr_routed_sdr', 'segr_routing_rule', 'headcount_source', 'segr_evidence (first signal)', 'segr_run_id']]
    for a in A:
        s = a['score']['signals']; ev = f'{s[0]["signal_type"]}: "{s[0]["evidence_quote"][:140]}" {s[0].get("source_url", "")}' if s else ', '.join(a['score'].get('guard_flags') or [])
        co.append([a['domain'], a['company'], a['segment'], a['tier'], a['score']['fit'], a['score']['intent'], a['score']['lane'] == 'needs_human', a['sdr'].get('sdr_name') or a['sdr']['queue'], a['sdr']['rule'], a.get('headcount_source', ''), ev, a['run_id']])
    ow = [['sdr_id', 'name (fictional)', 'territories', 'segments', 'capacity', 'assigned_this_run', 'hubspot_owner_email (to map)']] + [[r['sdr_id'], r['name'], ', '.join(r['territories']), ', '.join(r['segments']), r['capacity'], sum(1 for a in A if a['sdr'].get('sdr_id') == r['sdr_id']), ''] for r in R]
    tk = [['company (association 192)', 'hs_task_subject', 'hs_task_priority', 'hs_task_status', 'hs_timestamp (due, UTC)', 'owner (routed SDR)']]
    for a in A:
        if H.wants_task(a): p = H.task_props(a, None, now); tk.append([a['domain'], p['hs_task_subject'], p['hs_task_priority'], p['hs_task_status'], p['hs_timestamp'], a['sdr']['sdr_name']])
    wb = openpyxl.Workbook(); first = True
    for name, rows, note in [('Companies', co, 'HubSpot Companies view (secondary). Same records router/hubspot.py upserts, keyed on segr_domain_key. SDR names are fictional.'), ('Owners', ow, 'SDR roster (fictional).'), ('Tasks', tk, 'One task per Tier 1/2 routed account; reruns update the same task.')]:
        ws = wb.active if first else wb.create_sheet(); first = False; ws.title = name; ws.append([note]); ws['A1'].font = Font(italic=True)
        for r in rows: ws.append(r)
        for c in ws[2]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='D9E2F3')
        ws.freeze_panes = 'B3'
    wb.save(out_path)
if __name__ == '__main__': main(sys.argv[1], sys.argv[2])
